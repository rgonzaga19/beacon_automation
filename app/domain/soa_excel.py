"""SOA Excel generation extracted from the standalone generator."""

from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import re
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


TEMPLATE_PATH = Path(__file__).resolve().parents[2] / "templates" / "soa_master.xlsx"
CLAIM_FILLS = (
    PatternFill("solid", fgColor="FFFFF8E7"),
    PatternFill("solid", fgColor="FFEAF3FF"),
)

STATIC_DRUG_ROWS = [
    ("MED01", "APRINOL", "REGULAR HEPARIN", 1, 250, "VIAL", "DURING HD TREATMENT", "IV", "HEPARIN (as SODIUM) 5000 IU/mL SOLUTION 5 mL VIAL", "VIAL"),
    ("MED02", "SALINE PHILRX", "PNSS(PLAIN NORMAL SALINE SOLUTION)", 1, 100, "BOTTLE", "DURING HD TREATMENT", "IV", "0.9% SODIUM CHLORIDE SOLUTION 1 L BOTTLE", "BOTTLE"),
    ("MED03", "CITRAPURE", "HEMODIALYSIS ACID CONCENTRATE (DIALYSATE ACETATE BASED)", 1, 375, "GALLON", "DURING HD TREATMENT", "DIALYSATE", "HEMODIALYSIS ACID CONCENTRATE (DIALYSATE ACETATE BASED) 5 L", "GALLON"),
    ("MED04", "RENAL PURE", "HEMODIALYSIS BICARBONATE CONCENTRATE", 1, 175, "GALLON", "DURING HD TREATMENT", "DIALYSATE", "HEMODIALYSIS BICARBONATE CONCENTRATE 5 L", "GALLON"),
]

LAB_ROWS = [
    ("LABORATORY", "Chemistry - Sodium", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: SODIUM"),
    ("LABORATORY", "Chemistry - Potassium", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: POTASSIUM"),
    ("LABORATORY", "Chemistry - Serum Calcium", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: SERUM CALCIUM"),
    ("LABORATORY", "Chemistry - BUN", 1, 250, "LaboratoryAndDiagnostic", "", "CHEMISTRY: BUN"),
    ("LABORATORY", "Chemistry - URIC ACID", 1, 250, "LaboratoryAndDiagnostic", "", "CHEMISTRY: URIC ACID"),
    ("LABORATORY", "Chemistry - Creatinine", 1, 200, "LaboratoryAndDiagnostic", "", "CHEMISTRY: CREATININE"),
    ("LABORATORY", "Chemistry - UREAN NITROGEN", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: URINE-UREAN NITROGEN"),
    ("LABORATORY", "Chemistry - Serum Albumin", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: SERUM ALBUMIN"),
    ("LABORATORY", "Chemistry - Phosphorous", 1, 175, "LaboratoryAndDiagnostic", "", "CHEMISTRY: PHOSPHOROUS"),
    ("LABORATORY", "Hematology - CBC", 1, 250, "LaboratoryAndDiagnostic", "", "HEMATOLOGY: CBC"),
]


def _room_board_prices(claim):
    has_lab = bool(claim.get("hasLab"))
    has_epo = bool(claim.get("hasEpo"))
    epo_type = str(claim.get("epoType") or "alfa").lower()
    epo_qty = min(int(claim.get("epoQty") or 0), 1) if epo_type == "beta" else int(claim.get("epoQty") or 0)
    prices = {
        (False, False): (1750, 1750, 1162.5),
        (False, True): (950, 950, 762.5),
        (True, "alfa", 1, False): (1500, 1500, 787.5),
        (True, "alfa", 1, True): (700, 700, 387.5),
        (True, "beta", 1, False): (1250, 1250, 662.5),
        (True, "beta", 1, True): (450, 450, 262.5),
        (True, "alfa", 2, False): (1250, 1250, 412.5),
        (True, "alfa", 2, True): (350, 350, 212.5),
    }
    key = (has_epo, epo_type, epo_qty, has_lab) if has_epo else (False, has_lab)
    values = prices.get(key)
    if values is None:
        values = (1750, 1750, 1162.5)
    return (500, *values)


def _date_value(value):
    try:
        return datetime.fromisoformat(str(value)).date()
    except (TypeError, ValueError):
        raise ValueError("Each claim needs a valid renderDate in YYYY-MM-DD format.")


def _fill_rows(sheet, start, count, columns, claim_index):
    fill = CLAIM_FILLS[claim_index % len(CLAIM_FILLS)]
    for row in sheet.iter_rows(min_row=start, max_row=start + count - 1, min_col=1, max_col=columns):
        for cell in row:
            cell.fill = fill


def _insert_sheet_rows(sheet, rows, render_date, claim_index, number_columns):
    start = sheet.max_row + 1
    for values in rows:
        sheet.append([*values[:-1], render_date, values[-1]])
    _fill_rows(sheet, start, len(rows), number_columns, claim_index)
    for row in range(start, start + len(rows)):
        if number_columns == 12:
            sheet.cell(row, 5).number_format = "0.00"
            sheet.cell(row, 6).number_format = "0.00"
            sheet.cell(row, 8).number_format = "m/d/yyyy"
        else:
            sheet.cell(row, 4).number_format = "0.00"
            sheet.cell(row, 5).number_format = "0.00"
            sheet.cell(row, 6).number_format = "m/d/yyyy"


def _drug_rows(claim):
    rows = list(STATIC_DRUG_ROWS)
    if claim.get("hasEpo"):
        beta = str(claim.get("epoType") or "alfa").lower() == "beta"
        qty = min(int(claim.get("epoQty") or 1), 1) if beta else int(claim.get("epoQty") or 1)
        rows.append(("MED05", "RECORMON" if beta else "EPOETINE", "EPOETIN BETA" if beta else "EPOETIN ALFA", qty, 1500 if beta else 875, "PREFILLED SYRINGE", "POST HD TREATMENT", "IV/SUBCUTANEOUS", "EPOETIN BETA (RECOMBINANT ERYTHROPOIETIN) 5000IU/0.3ml SOLUTION PRE-FILLED SYRINGE WITH NEEDLE" if beta else "EPOETIN ALFA (RECOMBINANT HUMAN ERYTHROPOIETIN) 4000 IU/mL SOLUTION 1 mL PRE-FILLED GLASS SYRINGE", "VIAL"))
    return rows


def _supply_rows(data, claim):
    subkit = data.get("accessType") == "subkit"
    low_flux = data.get("fluxType") == "low"
    rows = [
        ("Supplies", "DIALYZER", 1, 500, "MedicalSupplies", "PIECE", "LOW FLUX DIALYZER OR ITS EQUIVALENT" if low_flux else "DIALYZER, CELLULOSE HIGH EFFICIENCY AM-SD-750U"),
        ("Supplies", "BLOODLINES", 1, 312.5, "MedicalSupplies", "PIECE", "BLOODLINES"),
        ("Supplies" if subkit else "Others", "Sterile Dressing Kit" if subkit else "Fistula Needles with Safety", 1, 312.5, "MedicalSupplies" if subkit else "Others", "KIT" if subkit else "", "DRESSING KIT" if subkit else "Others"),
        ("Supplies", "SUBCLAVIAN KIT" if subkit else "FISTULA KIT", 1, 312.5, "MedicalSupplies", "KIT", "IJ CATHETER FR. 12" if subkit else "FISTULA KIT"),
    ]
    room = _room_board_prices(claim)
    for description, price in zip(("Use of Machines", "Personnel Costs", "Rentals and Utilities", "Other Administrative Costs"), room):
        rows.append(("Others", f"Room and Board - {description}", 1, price, "RoomAndBoard", "", "ROOM AND BOARD"))
    if claim.get("hasLab"):
        rows.extend(LAB_ROWS)
    return rows


def generate_workbook(data, *, validate_epo_quantity=False, validate_laboratory=False):
    claims = data.get("claims") if isinstance(data, dict) else None
    if not claims or len(claims) > 7:
        raise ValueError("Provide between 1 and 7 claims.")
    if validate_laboratory and sum(bool(claim.get("hasLab")) for claim in claims) > 1:
        raise ValueError("Laboratory can only be included in one claim. Uncheck it in the other claims.")
    if validate_epo_quantity:
        for index, claim in enumerate(claims, 1):
            if not claim.get("hasEpo"):
                continue
            epo_type = str(claim.get("epoType") or "alfa").lower()
            maximum = 1 if epo_type == "beta" else 2
            quantity = claim.get("epoQty")
            if isinstance(quantity, bool) or not isinstance(quantity, (int, float)) or quantity not in range(1, maximum + 1):
                raise ValueError(f"Claim {index}: EPO {epo_type.title()} quantity must be a whole number from 1 to {maximum}.")
    workbook = load_workbook(TEMPLATE_PATH)
    drugs = workbook.worksheets[0]
    supplies = workbook.worksheets[1]
    if drugs.max_row > 2:
        drugs.delete_rows(3, drugs.max_row - 2)
    if supplies.max_row > 2:
        supplies.delete_rows(3, supplies.max_row - 2)
    for index, claim in enumerate(claims):
        render_date = _date_value(claim.get("renderDate"))
        drug_values = [(a, b, c, d, e, d * e, f, g, h, i, j) for a, b, c, d, e, f, g, h, i, j in _drug_rows(claim)]
        start = drugs.max_row + 1
        for values in drug_values:
            # Insert the render date before frequency, preserving all 12 columns.
            drugs.append(values[:7] + (render_date,) + values[7:])
        _fill_rows(drugs, start, len(drug_values), 12, index)
        for row in range(start, start + len(drug_values)):
            drugs.cell(row, 5).number_format = drugs.cell(row, 6).number_format = "0.00"
            drugs.cell(row, 8).number_format = "m/d/yyyy"
        supply_values = [(a, b, c, d, c * d, render_date, e, f, g) for a, b, c, d, e, f, g in _supply_rows(data, claim)]
        start = supplies.max_row + 1
        for values in supply_values:
            supplies.append(values)
        _fill_rows(supplies, start, len(supply_values), 9, index)
        for row in range(start, start + len(supply_values)):
            supplies.cell(row, 4).number_format = supplies.cell(row, 5).number_format = "0.00"
            supplies.cell(row, 6).number_format = "m/d/yyyy"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def build_batch_template():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Batch"
    for index, width in enumerate((34, 20, 20, 20, 20, 20, 20, 20), 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    headers = [
        "NAME OF PATIENT",
        "TREATMENT DATES",
        "NO. OF CLAIMS",
        "DATES OF ERYTHROPOIETIN GIVEN\n(WEEKLY)",
        "BETA RECORMON",
        "DIALYZER CATEGORY",
        "KIT CATEGORY",
        "W/ LAB",
    ]
    sheet.append(headers)
    sheet.row_dimensions[1].height = 48
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)

    sample_rows = [
        ["ABOGADO", "Jul 28,30", 2, "28,30", "", "FISTULA", "HIGH FLUX", "28"],
        ["ALARZAR", "Jul 27,29,31", 3, "27,31", "", "SUBKIT", "LOW FLUX", "27"],
        ["EXAMPLE 1 - dose qty (2)", "Jul 27,31", 2, "27(2),31", "", "FISTULA", "HIGH FLUX", "NO"],
        ["EXAMPLE 2 - dose qty x2", "Jul 27,31", 2, "27x2,31", "", "SUBKIT", "LOW FLUX", "NO"],
        ["EXAMPLE 3 - repeated day = qty 2", "Jul 27,31", 2, "27,27,31,31", "", "SUBKIT", "HIGH FLUX", "NO"],
    ]
    for values in sample_rows:
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.alignment = Alignment(vertical="center", horizontal="center")
    sheet.freeze_panes = "A2"

    notes = workbook.create_sheet("Instructions")
    notes.column_dimensions["A"].width = 40
    notes.column_dimensions["B"].width = 100
    notes.append(["How to fill out the Batch sheet"])
    notes["A1"].font = Font(bold=True, size=13)
    notes.append([])

    note_lines = [
        ["Non-blocking review", "Upload the workbook to see row-level warnings. Warnings do not block generation. Skipped rows, omitted dates, and adjusted values are listed in Validation_Report.txt inside the ZIP when warnings exist."],
        ["Date separators", "Periods are also accepted as day separators: 5.7.9 means days 5, 7, and 9. This applies to treatment, EPO, Beta, and lab days. Use a named month or an Excel date for complete dates; dotted numeric dates are interpreted as day lists."],
        ["NAME OF PATIENT", "Used as the output file name."],
        ["TREATMENT DATES", "Accepts a real date, \"29-Jul\", \"Jul 28,30\", a day range (\"Jul 27-29\"), ordinal days (\"29th Jul\", \"1st,3rd\"), or an explicit year anywhere in the text (\"Jul 28, 2027\") to override the app's Default Claim Period year. Days can be separated with commas, semicolons, or slashes (\"5;7;9\", \"5/7/9\"). If you only type day numbers with no month (e.g. \"28,30\" or \"27-29\"), the app's Default Claim Period month/year fills the gap."],
        ["NO. OF CLAIMS", "Informational - the app counts claims from Treatment Dates directly. A mismatch just shows as a warning, it won't block generation."],
        ["DATES OF ERYTHROPOIETIN GIVEN (WEEKLY)", "Day number(s) EPO Alfa was given - must match a day already listed in Treatment Dates. A second dose on the same day: write \"27(2)\", \"27x2\", \"27*2\", or list the day twice (\"27,27\"). A range like \"27-29\" applies to every day in it; add a quantity to the whole range with \"27-29x2\". Ordinal suffixes (\"27th\") and semicolon/slash separators are also accepted. Max 2 doses per day."],
        ["BETA RECORMON", "Same formats and matching rules as Erythropoietin above, but max 1 dose per day."],
        ["W/ LAB", "Day number(s) a lab was included - accepts the same day-range, ordinal, and separator formats as the other columns - or \"NO\" / blank for none."],
        ["DIALYZER CATEGORY", "FISTULA or SUBKIT."],
        ["KIT CATEGORY", "HIGH FLUX or LOW FLUX."],
    ]
    for label, description in note_lines:
        notes.append([label, description])
        row = notes[notes.max_row]
        row[0].font = Font(bold=True)
        row[0].alignment = Alignment(vertical="top")
        row[1].alignment = Alignment(vertical="top", wrap_text=True)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _normalized_row(row):
    return {re.sub(r"[^A-Z0-9]+", " ", str(key).upper()).strip(): value for key, value in row.items() if key is not None}


def _days(value, warn=None):
    if value is None or str(value).strip().lower() in ("", "no", "none", "n/a", "-"):
        return {}
    if isinstance(value, (date, datetime)):
        return {value.day: 1}
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    result = {}
    text = re.sub(r"(\d)(st|nd|rd|th)", r"\1", str(value), flags=re.I).replace(" ", "")
    for token in re.split(r"[.,;/\n]+", text):
        if not token:
            continue
        token = token.strip("-")
        match = re.fullmatch(r"(\d{1,2})(?:-(\d{1,2}))?(?:[x*(](\d+)\)?)?", token)
        if not match:
            if warn:
                warn(f"Unrecognized entry '{token}' was ignored.")
            continue
        start, end, quantity = int(match.group(1)), int(match.group(2) or match.group(1)), int(match.group(3) or 1)
        if not 1 <= start <= end <= 31 or quantity < 1:
            if warn:
                warn(f"Invalid day range or quantity '{token}' was ignored.")
            continue
        for day in range(start, end + 1):
            if 1 <= day <= 31:
                result[day] = result.get(day, 0) + quantity
    return result


def _treatment(value, default_month, default_year, warn=None):
    if isinstance(value, (date, datetime)):
        return value.day, value.month, value.year
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"(\d)(st|nd|rd|th)", r"\1", str(value or "").strip(), flags=re.I)
    months = {name[:3].lower(): index for index, name in enumerate(("January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"), 1)}
    month_match = re.search(r"([A-Za-z]{3,9})", text)
    month = months.get(month_match.group(1)[:3].lower()) if month_match else default_month
    year_match = re.search(r"(?:^|\D)(20\d{2}|19\d{2})(?:\D|$)", text)
    year = int(year_match.group(1)) if year_match else default_year
    text = re.sub(r"(?<!\d)(?:20\d{2}|19\d{2})(?!\d)", "", text)
    if month is None:
        if warn:
            warn("Unrecognized month; the default claim month was used.")
        month = default_month
    if month_match:
        day_text = text.replace(month_match.group(1), "")
        days = _days(day_text, warn)
    else:
        days = _days(text, warn)
    return (next(iter(days)), month, year) if len(days) == 1 else (days, month, year)


def batch_workbooks(file_stream, month, year, *, preview=False):
    source = load_workbook(file_stream, data_only=True).active
    output = BytesIO()
    generated = 0
    warnings = []
    def warn(row, patient, column, message):
        warnings.append({"row": row, "patient": patient, "column": column, "message": message})
    with ZipFile(output, "w", ZIP_DEFLATED) as archive:
        names = {}
        headers = [cell.value for cell in source[1]]
        expected = ("NAME OF PATIENT", "TREATMENT DATES", "NO OF CLAIMS", "DATES OF ERYTHROPOIETIN GIVEN WEEKLY", "BETA RECORMON", "DIALYZER CATEGORY", "KIT CATEGORY", "W LAB")
        normalized_headers = _normalized_row(dict.fromkeys(headers))
        for column in expected:
            if column not in normalized_headers:
                warn(1, "", column, "Column is missing. Its values will be treated as blank.")
        for row_number, values in enumerate(source.iter_rows(min_row=2, values_only=True), 2):
            if all(value is None or str(value).strip() == "" for value in values):
                continue
            row = _normalized_row(dict(zip(headers, values)))
            name = str(row.get("NAME OF PATIENT") or "Unnamed").strip() or "Unnamed"
            def issue(column, message):
                warn(row_number, name, column, message)
            if not str(row.get("NAME OF PATIENT") or "").strip():
                issue("NAME OF PATIENT", "Missing patient name; Unnamed will be used.")
            treatment_result = _treatment(row.get("TREATMENT DATES"), month, year, lambda message: issue("TREATMENT DATES", message))
            if not treatment_result:
                continue
            treatment_days, claim_month, claim_year = treatment_result
            if isinstance(treatment_days, int):
                treatment_days = {treatment_days: 1}
            for day in list(treatment_days):
                try:
                    date(claim_year, claim_month, day)
                except ValueError:
                    issue("TREATMENT DATES", f"Invalid date {claim_year}-{claim_month:02d}-{day:02d} was omitted.")
                    del treatment_days[day]
            declared = row.get("NO OF CLAIMS")
            if declared is not None and str(declared).strip() and str(declared).strip() != str(len(treatment_days)):
                issue("NO. OF CLAIMS", f"Entered count {declared} differs from {len(treatment_days)} usable treatment dates; dates determine the count.")
            if len(treatment_days) > 7:
                issue("TREATMENT DATES", "Only the first 7 usable treatment dates will be included; remaining dates are omitted.")
            access = "subkit" if str(row.get("DIALYZER CATEGORY") or "").strip().lower() == "subkit" else "fistula"
            flux = "low" if "low" in str(row.get("KIT CATEGORY") or "").lower() else "high"
            if str(row.get("DIALYZER CATEGORY") or "").strip().lower() not in ("subkit", "fistula"):
                issue("DIALYZER CATEGORY", "Missing or unknown category; FISTULA will be used.")
            if str(row.get("KIT CATEGORY") or "").strip().lower() not in ("high flux", "low flux"):
                issue("KIT CATEGORY", f"Missing or unknown category; {flux.upper()} FLUX will be used.")
            alfa = _days(row.get("DATES OF ERYTHROPOIETIN GIVEN WEEKLY"), lambda message: issue("EPO ALFA", message))
            beta = _days(row.get("BETA RECORMON"), lambda message: issue("BETA RECORMON", message))
            labs = set(_days(row.get("W LAB"), lambda message: issue("W/ LAB", message)))
            included_days = set(list(treatment_days)[:7])
            for column, days in (("EPO ALFA", alfa), ("BETA RECORMON", beta), ("W/ LAB", labs)):
                unmatched = set(days) - included_days
                if unmatched:
                    issue(column, f"Days {', '.join(map(str, sorted(unmatched)))} have no included treatment date and will be ignored.")
            for day, quantity in alfa.items():
                if quantity > 2:
                    issue("EPO ALFA", f"Day {day} has {quantity} doses (expected at most 2); the entered quantity is retained. Review before use.")
            for day, quantity in beta.items():
                if quantity > 1:
                    issue("BETA RECORMON", f"Day {day} has {quantity} doses; generation limits this to 1.")
            for day in alfa.keys() & beta.keys():
                issue("EPO", f"Day {day} lists both Alfa and Beta; Beta will be used.")
            claims = []
            for day, _ in list(treatment_days.items())[:7]:
                is_beta = day in beta
                has_epo = is_beta or day in alfa
                claims.append({"renderDate": f"{claim_year:04d}-{claim_month:02d}-{day:02d}", "hasEpo": has_epo, "epoQty": min(beta.get(day, 1), 1) if is_beta else alfa.get(day, 1), "epoType": "beta" if is_beta else "alfa", "hasLab": day in labs})
            if not claims:
                issue("TREATMENT DATES", "No usable treatment dates; this patient row will be skipped.")
                continue
            safe_name = re.sub(r"[\\/:*?\"<>|]", "_", name)[:100] or "Unnamed"
            count = names.get(safe_name, 0) + 1
            names[safe_name] = count
            filename = f"{safe_name}{f'_{count}' if count > 1 else ''}.xlsx"
            if not preview:
                workbook = generate_workbook({"accessType": access, "fluxType": flux, "claims": claims})
                archive.writestr(filename, workbook.getvalue())
            generated += 1
        if warnings and not preview:
            report = [f"Generated {generated} SOA files. Review warnings before use.", ""]
            report.extend(f"Row {item['row']} | {item['patient']} | {item['column']}: {item['message']}" for item in warnings)
            archive.writestr("Validation_Report.txt", "\n".join(report))
    if preview:
        return {"warnings": warnings, "generated": generated}
    if not generated:
        raise ValueError("No usable treatment dates were found in the workbook.")
    output.seek(0)
    return output
